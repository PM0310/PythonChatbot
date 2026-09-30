"""
Sales Order Chatbot — Flask backend with custom UI (v2 – smart create order).
Calls Oracle EBS REST API: GET_ORDER to retrieve sales order details.
"""

import json
import ssl
import re
import datetime
import smtplib
import threading
import time
from email.mime.text import MIMEText
from email.mime.multipart import MIMEMultipart
import requests
import requests.adapters
import urllib3
from urllib3.util.ssl_ import create_urllib3_context
from requests.auth import HTTPBasicAuth
from flask import Flask, request, jsonify, send_from_directory, Response

# ── Optional voice packages ───────────────────────
try:
    from vosk import Model, KaldiRecognizer
    import sounddevice as sd
    VOSK_AVAILABLE = True
except ImportError:
    VOSK_AVAILABLE = False

try:
    import pyttsx3
    PYTTSX3_AVAILABLE = True
except ImportError:
    PYTTSX3_AVAILABLE = False

try:
    from word2number import w2n
    W2N_AVAILABLE = True
except ImportError:
    W2N_AVAILABLE = False

try:
    from openai import OpenAI
    OPENAI_AVAILABLE = True
except ImportError:
    OPENAI_AVAILABLE = False

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False

try:
    import oracledb
    ORACLEDB_AVAILABLE = True
except ImportError:
    ORACLEDB_AVAILABLE = False


# ── Suppress only the InsecureRequestWarning ──────
urllib3.disable_warnings(urllib3.exceptions.InsecureRequestWarning)

# ── Configuration ─────────────────────────────────
BASE_URL = "https://cendb.ad.centroid.com:4463"
GET_ORDER_URL = f"{BASE_URL}/webservices/rest/sales_order/GET_ORDER/"
PROCESS_ORDER_URL = f"{BASE_URL}/webservices/rest/sales_order/PROCESS_ORDER/"
CUST_SO_DTLS_URL = f"{BASE_URL}/webservices/rest/CustSODtls/XX_CUST_SO_DET_PRC1/"
REST_USERNAME = "operations"
REST_PASSWORD = "welcome"
ORG_ID = "204"
HEADERS = {"Content-Type": "application/json", "Accept": "application/json"}

# ── Oracle DB configuration (for lookup queries) ──
DB_USER = "apps"
DB_PASSWORD = "apps"
DB_DSN = "(DESCRIPTION=(ADDRESS=(PROTOCOL=tcp)(HOST=cendb.centroid.com)(PORT=1541))(CONNECT_DATA=(SERVICE_NAME=EBS122)))"
ORDER_TYPE_ID = 1437
OPERATING_UNIT = "Vision Operations"
MAX_CHOICES = 10


def _get_db_connection():
    """Get an oracledb connection (thick mode for legacy password verifiers)."""
    if not ORACLEDB_AVAILABLE:
        raise RuntimeError("oracledb is not installed. Run: pip install oracledb")
    try:
        oracledb.init_oracle_client()
    except Exception:
        pass  # already initialized
    return oracledb.connect(user=DB_USER, password=DB_PASSWORD, dsn=DB_DSN)

# ── AI Summary configuration ─────────────────────
GITHUB_TOKEN = "REDACTED_SECRET"
GITHUB_MODEL = "gpt-4o"


def generate_ai_summary(operation_type: str, data: dict) -> str:
    """Generate a concise AI summary of an ERP operation result using GPT-4o."""
    if not OPENAI_AVAILABLE:
        return ""

    try:
        client = OpenAI(
            base_url="https://models.inference.ai.azure.com",
            api_key=GITHUB_TOKEN,
        )

        today = datetime.datetime.now().strftime("%B %d, %Y")
        data_json = json.dumps(data, indent=2, default=str)[:3000]  # truncate for token limit

        system_prompt = (
            f"You are an ERP Sales Order Analyst. Today is {today}. "
            "Given the result of an ERP operation, write a concise executive summary (3-5 sentences). "
            "Highlight key details: order number, customer, status, amounts, line items count, "
            "any issues or exceptions. Be professional and actionable. "
            "Do NOT use markdown formatting — plain text only."
        )

        if operation_type == "order":
            user_prompt = f"Summarize this sales order lookup result:\n{data_json}"
        elif operation_type == "customer_orders":
            user_prompt = f"Summarize these customer order details:\n{data_json}"
        elif operation_type == "create":
            user_prompt = f"Summarize this order creation result:\n{data_json}"
        elif operation_type == "orders_advanced":
            user_prompt = f"Summarize this orders report (filtered query result):\n{data_json}"
        elif operation_type == "add_line":
            user_prompt = f"Summarize this add-line-to-order result:\n{data_json}"
        else:
            user_prompt = f"Summarize this ERP operation result:\n{data_json}"

        response = client.chat.completions.create(
            model=GITHUB_MODEL,
            messages=[
                {"role": "system", "content": system_prompt},
                {"role": "user", "content": user_prompt},
            ],
            temperature=0.3,
            max_tokens=300,
        )
        return response.choices[0].message.content.strip()
    except Exception as e:
        print(f"AI summary error: {e}")
        return ""

# ── Voice configuration ──────────────────────────
VOSK_MODEL_PATH = r"C:\Users\pavan.maccha\Desktop\Python\Model\vosk-model-small-en-us-0.15"
SAMPLE_RATE = 16000
VOICE_DURATION = 5  # seconds to record
_vosk_model = None


def _load_vosk_model():
    """Lazy-load the Vosk model."""
    global _vosk_model
    if _vosk_model is None and VOSK_AVAILABLE:
        import os
        if os.path.exists(VOSK_MODEL_PATH):
            _vosk_model = Model(VOSK_MODEL_PATH)
            print(f"Vosk model loaded from {VOSK_MODEL_PATH}")
        else:
            print(f"WARNING: Vosk model not found at {VOSK_MODEL_PATH}")
    return _vosk_model


def recognize_speech(duration: int = VOICE_DURATION) -> str:
    """Record audio for `duration` seconds and return recognized text using Vosk."""
    model = _load_vosk_model()
    if model is None:
        return ""
    rec = KaldiRecognizer(model, SAMPLE_RATE)
    print(f"\U0001f399  Listening for {duration} seconds...")
    audio = sd.rec(int(duration * SAMPLE_RATE), samplerate=SAMPLE_RATE,
                   channels=1, dtype="int16")
    sd.wait()
    if rec.AcceptWaveform(audio.tobytes()):
        result = json.loads(rec.Result())
    else:
        result = json.loads(rec.FinalResult())
    text = result.get("text", "").strip()
    print(f"\U0001f4dd  Recognised: '{text}'")
    return text


_tts_lock = threading.Lock()


def speak(text: str):
    """Speak text aloud using pyttsx3 (thread-safe)."""
    if not PYTTSX3_AVAILABLE:
        return
    print(f"\U0001f50a  {text}")
    with _tts_lock:
        engine = pyttsx3.init()
        engine.setProperty("rate", 160)
        engine.say(text)
        engine.runAndWait()
        engine.stop()


# ── Voice text normalization ──────────────────────
_WORD_DIGITS = {
    "zero": "0", "one": "1", "two": "2", "three": "3", "four": "4",
    "five": "5", "six": "6", "seven": "7", "eight": "8", "nine": "9",
}

_NUMBER_WORDS = {
    "zero", "one", "two", "three", "four", "five", "six", "seven",
    "eight", "nine", "ten", "eleven", "twelve", "thirteen", "fourteen",
    "fifteen", "sixteen", "seventeen", "eighteen", "nineteen", "twenty",
    "thirty", "forty", "fifty", "sixty", "seventy", "eighty", "ninety",
    "hundred", "thousand", "million", "billion",
}


def _convert_number_words(words: list) -> str:
    """Convert a list of number-word tokens to a digit string.

    Strategy:
      1. Try w2n on the whole phrase (handles 'sixty seven thousand' well)
      2. If w2n fails or gives implausible result, split into sub-groups
         separated by single-digit words and convert each group
    """
    if not words:
        return ""

    phrase = " ".join(words)

    # Check if the sequence contains multiplier words (hundred/thousand etc.)
    # which indicate a proper compound number that w2n handles well
    compound = {"ten", "eleven", "twelve", "thirteen", "fourteen", "fifteen",
                "sixteen", "seventeen", "eighteen", "nineteen", "twenty",
                "thirty", "forty", "fifty", "sixty", "seventy", "eighty",
                "ninety", "hundred", "thousand", "million", "billion"}
    tens_words = {"twenty", "thirty", "forty", "fifty", "sixty", "seventy",
                  "eighty", "ninety"}
    multipliers = {"hundred", "thousand", "million", "billion"}
    has_multiplier = any(w.lower() in multipliers for w in words)
    all_single = all(w.lower() in _WORD_DIGITS for w in words)

    # Strategy 1: try w2n on whole phrase if multiplier words are present
    # (e.g. "sixty seven thousand nine hundred sixty five")
    # OR if all words form a valid compound without standalone digits
    if W2N_AVAILABLE and (has_multiplier or (not all_single and len(words) <= 3)):
        try:
            result = str(w2n.word_to_num(phrase))
            return result
        except Exception:
            pass

    # Strategy 2: smart split — group compound words with their following
    # single-digit (e.g. "sixty seven" → 67), standalone singles as digits
    result_parts = []
    compound_buf = []

    def flush_compound():
        if not compound_buf:
            return
        sub = " ".join(compound_buf)
        if W2N_AVAILABLE:
            try:
                result_parts.append(str(w2n.word_to_num(sub)))
            except Exception:
                for w in compound_buf:
                    result_parts.append(_WORD_DIGITS.get(w.lower(), w))
        else:
            for w in compound_buf:
                result_parts.append(_WORD_DIGITS.get(w.lower(), w))
        compound_buf.clear()

    for w in words:
        wl = w.lower()
        if wl in compound:
            compound_buf.append(w)
        elif wl in _WORD_DIGITS:
            # "seven" after "sixty" → part of compound (sixty seven = 67)
            if compound_buf and compound_buf[-1].lower() in tens_words:
                compound_buf.append(w)
            else:
                flush_compound()
                result_parts.append(_WORD_DIGITS[wl])
        else:
            flush_compound()
            result_parts.append(w)

    flush_compound()
    return "".join(result_parts)


def normalize_voice_text(text: str) -> str:
    """Convert spoken text to a form the intent matcher can handle.

    - Converts spoken numbers like 'sixty seven nine sixty five' → '67965'
    - Handles digit-by-digit: 'six seven nine six five' → '67965'
    - Preserves non-numeric words as-is
    """
    if not text:
        return text

    text = text.strip()

    tokens = text.split()
    result_tokens = []
    num_buffer = []

    def flush_number_buffer():
        if not num_buffer:
            return
        converted = _convert_number_words(num_buffer)
        result_tokens.append(converted)
        num_buffer.clear()

    for token in tokens:
        if token.lower() in _NUMBER_WORDS:
            num_buffer.append(token)
        else:
            flush_number_buffer()
            result_tokens.append(token)

    flush_number_buffer()
    return " ".join(result_tokens)


# ── TLS adapter for legacy Oracle EBS ciphers ────
class ForceTLSAdapter(requests.adapters.HTTPAdapter):
    def init_poolmanager(self, *args, **kwargs):
        ctx = create_urllib3_context()
        ctx.check_hostname = False
        ctx.verify_mode = ssl.CERT_NONE
        ctx.set_ciphers("DEFAULT@SECLEVEL=1")
        ctx.options |= ssl.OP_NO_SSLv2 | ssl.OP_NO_SSLv3
        kwargs["ssl_context"] = ctx
        return super().init_poolmanager(*args, **kwargs)


def get_order(order_number: int) -> dict:
    """Call Oracle EBS GET_ORDER REST service."""
    payload = {
        "PROCESS_ORDER_Input": {
            "RESTHeader": {
                "Responsibility": "ORDER_MGMT_SUPER_USER",
                "RespApplication": "ONT",
                "SecurityGroup": "STANDARD",
                "NLSLanguage": "AMERICAN",
                "Org_Id": ORG_ID,
            },
            "InputParameters": {
                "P_API_VERSION_NUMBER": 1,
                "P_INIT_MSG_LIST": "T",
                "P_RETURN_VALUES": "T",
                "P_ACTION_COMMIT": "T",
                "P_ORDER_NUMBER": order_number,
            },
        }
    }

    session = requests.Session()
    session.mount("https://", ForceTLSAdapter())
    session.auth = HTTPBasicAuth(REST_USERNAME, REST_PASSWORD)
    session.verify = False

    resp = session.post(
        GET_ORDER_URL,
        headers=HEADERS,
        data=json.dumps(payload),
        timeout=180,
    )

    if resp.status_code in (200, 201, 204):
        try:
            return resp.json()
        except Exception:
            return {"raw": resp.text}
    resp.raise_for_status()


def _build_session():
    """Create a requests session with TLS adapter and auth."""
    s = requests.Session()
    s.mount("https://", ForceTLSAdapter())
    s.auth = HTTPBasicAuth(REST_USERNAME, REST_PASSWORD)
    s.verify = False
    return s


def get_customer_orders(customer_name: str) -> dict:
    """Call CustSODtls REST service to get sales order details by customer name."""
    session = _build_session()
    resp = session.get(
        CUST_SO_DTLS_URL,
        headers=HEADERS,
        params={"P_CUSTOMER_NAME": customer_name},
        timeout=180,
    )
    if resp.status_code in (200, 201, 204):
        try:
            return resp.json()
        except Exception:
            return {"raw": resp.text}
    resp.raise_for_status()


def create_order(order_data: dict) -> dict:
    """Call Oracle EBS PROCESS_ORDER REST service to create a sales order."""
    payload = {
        "PROCESS_ORDER_Input": {
            "RESTHeader": {
                "Responsibility": "ORDER_MGMT_SUPER_USER",
                "RespApplication": "ONT",
                "SecurityGroup": "STANDARD",
                "NLSLanguage": "AMERICAN",
                "Org_Id": ORG_ID,
            },
            "InputParameters": {
                "P_API_VERSION_NUMBER": 1,
                "P_INIT_MSG_LIST": "T",
                "P_RETURN_VALUES": "T",
                "P_ACTION_COMMIT": "T",
                "P_HEADER_REC": {
                    "BOOKED_FLAG": "N",
                    "CUST_PO_NUMBER": order_data.get("cust_po_number", ""),
                    "ORDER_TYPE_ID": int(order_data.get("order_type_id", ORDER_TYPE_ID)),
                    "ORG_ID": int(ORG_ID),
                    "PAYMENT_TERM_ID": int(order_data.get("payment_term_id", 4)),
                    "PRICE_LIST_ID": int(order_data.get("price_list_id", 1000)),
                    "SALESREP_ID": int(order_data.get("salesrep_id", 10067)),
                    "SHIP_TO_ORG_ID": int(order_data.get("ship_to_org_id", 6472)),
                    "SOLD_TO_ORG_ID": int(order_data.get("sold_to_org_id", 5391)),
                    "TRANSACTIONAL_CURR_CODE": order_data.get("currency", "USD"),
                    "OPERATION": "CREATE",
                },
                "P_LINE_TBL": {
                    "P_LINE_TBL_ITEM": {
                        "INVENTORY_ITEM_ID": int(order_data.get("inventory_item_id", 12027)),
                        "ORDERED_ITEM": order_data.get("ordered_item", ""),
                        "ORDERED_QUANTITY": float(order_data.get("ordered_quantity", 1)),
                        "PAYMENT_TERM_ID": int(order_data.get("payment_term_id", 4)),
                        "PRICE_LIST_ID": int(order_data.get("price_list_id", 1000)),
                        "UNIT_LIST_PRICE": float(order_data.get("unit_list_price", 0)),
                        "UNIT_SELLING_PRICE": float(order_data.get("unit_selling_price", 0)),
                        "OPERATION": "CREATE",
                    }
                },
                "P_RTRIM_DATA": "n",
                "P_OPERATING_UNIT": OPERATING_UNIT,
                "P_DEBUG_LEVEL": 10,
            },
        }
    }

    session = _build_session()
    resp = session.post(
        PROCESS_ORDER_URL,
        headers=HEADERS,
        data=json.dumps(payload),
        timeout=180,
    )

    if resp.status_code in (200, 201, 204):
        try:
            return resp.json()
        except Exception:
            return {"raw": resp.text}
    resp.raise_for_status()


def get_orders_advanced(range_type=None, year=None, status=None):
    """Query orders by date range, year, or status (header-first, then line fallback)."""
    date_filter_dropped = False
    if year:
        date_clause = f"EXTRACT(YEAR FROM ooha.CREATION_DATE) = {int(year)}"
    elif range_type:
        date_map = {
            "this_month": "TRUNC(SYSDATE, 'MM')",
            "this_week": "TRUNC(SYSDATE, 'IW')",
            "last_30_days": "SYSDATE - 30",
        }
        date_clause = f"ooha.CREATION_DATE >= {date_map.get(range_type, 'SYSDATE - 7')}"
    else:
        date_clause = "1=1"

    bind_vars = {"org": int(ORG_ID)}
    matched_on = None  # tracks which column matched

    if status:
        status_upper = status.upper()
        bind_vars["status_val"] = f"%{status_upper}%"

        conn = _get_db_connection()
        cur = conn.cursor()

        # ── Step 1: Try header status first (with date filter) ──
        header_status_clause = "UPPER(ooha.FLOW_STATUS_CODE) LIKE :status_val"
        sql_count_hdr = (
            f"SELECT COUNT(DISTINCT ooha.ORDER_NUMBER) "
            f"FROM OE_ORDER_HEADERS_ALL ooha, OE_ORDER_LINES_ALL oola "
            f"WHERE ooha.HEADER_ID = oola.HEADER_ID "
            f"AND {date_clause} AND {header_status_clause} AND ooha.ORG_ID = :org"
        )
        print(f"🔍 Probe header status: {sql_count_hdr}")
        print(f"🔍 Bind vars: {bind_vars}")
        cur.execute(sql_count_hdr, bind_vars)
        hdr_count = cur.fetchone()[0]
        print(f"🔍 Header status count: {hdr_count}")

        if hdr_count > 0:
            matched_on = "header_status"
            status_clause = header_status_clause
        else:
            # ── Step 2: Try line status (with date filter) ──
            line_status_clause = "UPPER(oola.FLOW_STATUS_CODE) LIKE :status_val"
            sql_count_line = (
                f"SELECT COUNT(DISTINCT ooha.ORDER_NUMBER) "
                f"FROM OE_ORDER_HEADERS_ALL ooha, OE_ORDER_LINES_ALL oola "
                f"WHERE ooha.HEADER_ID = oola.HEADER_ID "
                f"AND {date_clause} AND {line_status_clause} AND ooha.ORG_ID = :org"
            )
            cur.execute(sql_count_line, bind_vars)
            line_count = cur.fetchone()[0]
            print(f"🔍 Line status count: {line_count}")

            if line_count > 0:
                matched_on = "line_status"
                status_clause = line_status_clause
            else:
                # ── Step 3: Try header status WITHOUT date filter ──
                sql_count_hdr_nodate = (
                    f"SELECT COUNT(DISTINCT ooha.ORDER_NUMBER) "
                    f"FROM OE_ORDER_HEADERS_ALL ooha, OE_ORDER_LINES_ALL oola "
                    f"WHERE ooha.HEADER_ID = oola.HEADER_ID "
                    f"AND {header_status_clause} AND ooha.ORG_ID = :org"
                )
                bind_nodate = {"org": int(ORG_ID), "status_val": f"%{status_upper}%"}
                cur.execute(sql_count_hdr_nodate, bind_nodate)
                hdr_nodate = cur.fetchone()[0]
                print(f"🔍 Header status (no date filter) count: {hdr_nodate}")

                if hdr_nodate > 0:
                    # Found in header but not for the given date range — drop date filter
                    matched_on = "header_status"
                    status_clause = header_status_clause
                    date_clause = "1=1"  # remove date filter to show results
                    date_filter_dropped = True
                    print(f"⚠️ No orders matched date+status. Dropping date filter — found {hdr_nodate} in header status.")
                else:
                    # ── Step 4: Try line status WITHOUT date filter ──
                    sql_count_line_nodate = (
                        f"SELECT COUNT(DISTINCT ooha.ORDER_NUMBER) "
                        f"FROM OE_ORDER_HEADERS_ALL ooha, OE_ORDER_LINES_ALL oola "
                        f"WHERE ooha.HEADER_ID = oola.HEADER_ID "
                        f"AND {line_status_clause} AND ooha.ORG_ID = :org"
                    )
                    cur.execute(sql_count_line_nodate, bind_nodate)
                    line_nodate = cur.fetchone()[0]
                    print(f"🔍 Line status (no date filter) count: {line_nodate}")

                    if line_nodate > 0:
                        matched_on = "line_status"
                        status_clause = line_status_clause
                        date_clause = "1=1"
                        date_filter_dropped = True
                        print(f"⚠️ No orders matched date+status. Dropping date filter — found {line_nodate} in line status.")
                    else:
                        matched_on = "none"
                        status_clause = "1=0"

        cur.close()
        conn.close()
    else:
        status_clause = "1=1"

    sql_count = (
        f"SELECT COUNT(DISTINCT ooha.ORDER_NUMBER) "
        f"FROM OE_ORDER_HEADERS_ALL ooha, OE_ORDER_LINES_ALL oola "
        f"WHERE ooha.HEADER_ID = oola.HEADER_ID "
        f"AND {date_clause} AND {status_clause} AND ooha.ORG_ID = :org"
    )
    sql_main = (
        f"SELECT ooha.ORDER_NUMBER, ooha.FLOW_STATUS_CODE AS HEADER_STATUS, "
        f"oola.FLOW_STATUS_CODE AS LINE_STATUS, "
        f"TO_CHAR(ooha.CREATION_DATE, 'DD-MON-YYYY') AS CREATION_DATE "
        f"FROM OE_ORDER_HEADERS_ALL ooha, OE_ORDER_LINES_ALL oola "
        f"WHERE ooha.HEADER_ID = oola.HEADER_ID "
        f"AND {date_clause} AND {status_clause} AND ooha.ORG_ID = :org "
        f"ORDER BY ooha.CREATION_DATE DESC, ooha.ORDER_NUMBER"
    )

    conn = _get_db_connection()
    cur = conn.cursor()
    cur.execute(sql_count, bind_vars)
    total_count = cur.fetchone()[0]
    cur.execute(sql_main, bind_vars)
    rows = cur.fetchall()
    cur.close()
    conn.close()
    print(f"✅ Final result: {total_count} orders, matched_on={matched_on}, date_clause={date_clause}")
    return {
        "total_count": total_count,
        "matched_on": matched_on,
        "date_filter_dropped": date_filter_dropped,
        "orders": [
            {"order_number": int(r[0]), "header_status": r[1] or "", "line_status": r[2] or "", "creation_date": r[3] or ""}
            for r in rows[:200]
        ],
    }


def add_line_payload(header_id, item_id, ordered_item, qty, price, selling_price, payment_term_id, price_list_id):
    """Build PROCESS_ORDER payload to add a line to an existing order."""
    return {
        "PROCESS_ORDER_Input": {
            "RESTHeader": {
                "Responsibility": "ORDER_MGMT_SUPER_USER", "RespApplication": "ONT",
                "SecurityGroup": "STANDARD", "NLSLanguage": "AMERICAN", "Org_Id": ORG_ID,
            },
            "InputParameters": {
                "P_API_VERSION_NUMBER": 1, "P_INIT_MSG_LIST": "T",
                "P_RETURN_VALUES": "T", "P_ACTION_COMMIT": "T",
                "P_HEADER_REC": {"HEADER_ID": int(header_id), "ORG_ID": int(ORG_ID), "OPERATION": "UPDATE"},
                "P_LINE_TBL": {
                    "P_LINE_TBL_ITEM": {
                        "HEADER_ID": int(header_id), "INVENTORY_ITEM_ID": int(item_id),
                        "ORDERED_ITEM": ordered_item, "ORDERED_QUANTITY": float(qty),
                        "PAYMENT_TERM_ID": int(payment_term_id), "PRICE_LIST_ID": int(price_list_id),
                        "UNIT_LIST_PRICE": float(price), "UNIT_SELLING_PRICE": float(selling_price),
                        "OPERATION": "CREATE",
                    }
                },
                "P_RTRIM_DATA": "n", "P_OPERATING_UNIT": OPERATING_UNIT, "P_DEBUG_LEVEL": 10,
            },
        }
    }


def ai_extract_intent(user_input: str) -> dict:
    """Use GPT-4o to extract structured intent from user input."""
    if not OPENAI_AVAILABLE:
        return {"action": "unknown"}
    try:
        client = OpenAI(
            base_url="https://models.inference.ai.azure.com",
            api_key=GITHUB_TOKEN,
        )
        current_date = datetime.datetime.now().strftime("%B %d, %Y")
        current_year = datetime.datetime.now().year
        system = (
            f"You are an ERP Sales Order Assistant. Today is {current_date} (Year: {current_year}).\n"
            "Extract the user's intent into JSON. Return ONLY raw JSON, no markdown.\n"
            "FIELDS:\n"
            "- action: 'create', 'get', 'get_range', 'add_line', 'email', 'help', 'general'\n"
            "- date_range: 'this_month', 'this_week', 'last_30_days', or null\n"
            "- year: integer (e.g. 2024) or null. If 'last year', compute from current year.\n"
            "- status: string ('BOOKED','ENTERED','CANCELLED','CLOSED','AWAITING_SHIPPING','BACKORDERED','SHIPPED','FULFILLED','AWAITING_FULFILLMENT','AWAITING_RETURN','PO_REQ_CREATED') or null\n"
            "- order_number: integer or null\n"
            "- customer_name: string or null\n"
            "- email_address: string or null\n"
            "- general_answer: string or null (if action='general', provide a helpful answer)\n"
            "- cust_po_number: string or null (Customer PO number, e.g. PO051, CPO123). Extract if action='create'.\n"
            "- ordered_item: string or null (item name/code, e.g. XP9004, AS54888). Extract if action='create'.\n"
            "- ordered_quantity: number or null (quantity, e.g. 85, 100). Extract if action='create'.\n"
            "\nRULES:\n"
            "1. If user mentions a time period (last 30 days, this month, this week, etc.) or a year, set action='get_range'.\n"
            "2. If user mentions a status (booked, entered, cancelled, closed, awaiting shipping, backordered, shipped, fulfilled, awaiting fulfillment), set action='get_range'. Convert spaces to underscores and uppercase for the status field.\n"
            "3. If user says create/new/place order, set action='create'. Also extract any order details provided inline (cust_po_number, ordered_item, ordered_quantity, customer_name).\n"
            "4. If user says add line/add item to order, set action='add_line'.\n"
            "5. If user provides an order number to look up, set action='get'.\n"
            "6. If user asks to send email, set action='email'.\n"
            "7. If user asks a general question (what is ERP, what is sales order, etc.), set action='general' and provide a concise answer in general_answer.\n"
            "8. If user provides a customer name for lookup, extract customer_name and set action='get'.\n"
            "9. If unsure, set action='general' and provide a helpful answer.\n"
        )
        response = client.chat.completions.create(
            model=GITHUB_MODEL,
            messages=[{"role": "system", "content": system}, {"role": "user", "content": user_input}],
            temperature=0.0,
        )
        clean = response.choices[0].message.content.strip().replace("```json", "").replace("```", "").strip()
        return json.loads(clean)
    except Exception as e:
        print(f"AI intent error: {e}")
        return {"action": "unknown"}


def parse_create_order_response(data: dict) -> dict:
    """Parse the PROCESS_ORDER response into a displayable summary."""
    try:
        # Try nested PROCESS_ORDER_Output path first (like reference file)
        output = data.get("PROCESS_ORDER_Output", {}).get("OutputParameters", {})
        if not output:
            output = data.get("OutputParameters", data)

        header = output.get("X_HEADER_REC", output.get("HEADER_REC", {}))
        return_status = output.get("X_RETURN_STATUS", output.get("RETURN_STATUS", ""))
        msg_count = output.get("X_MSG_COUNT", output.get("MSG_COUNT", 0))
        msg_data = output.get("X_MSG_DATA", output.get("MSG_DATA", ""))

        order_number = header.get("ORDER_NUMBER", "")
        header_id = header.get("HEADER_ID", "")
        status = header.get("FLOW_STATUS_CODE", "")
        cust_po = header.get("CUST_PO_NUMBER", "")

        # Fallback extraction
        if not order_number:
            out2 = data.get("OutputParameters", {})
            hdr2 = out2.get("X_HEADER_REC", {})
            order_number = hdr2.get("ORDER_NUMBER", "")
            header_id = header_id or hdr2.get("HEADER_ID", "")

        return {
            "order_number": order_number,
            "header_id": header_id,
            "return_status": return_status,
            "flow_status": status,
            "cust_po_number": cust_po,
            "msg_count": msg_count,
            "msg_data": msg_data,
            "success": return_status in ("S", "s"),
        }
    except Exception as e:
        return {"error": f"Failed to parse response: {e}", "raw": data}


def parse_customer_orders(data: dict) -> dict:
    """Parse the CustSODtls response into a displayable structure."""
    try:
        output = data.get("OutputParameters", data)
        # The result set can be nested under various keys
        rows_raw = (
            output.get("X_RESULT", None)
            or output.get("P_RESULT", None)
            or output.get("ResultSet", None)
            or output
        )
        # Normalise to list of dicts
        if isinstance(rows_raw, dict):
            inner = (
                rows_raw.get("X_RESULT_ITEM", None)
                or rows_raw.get("P_RESULT_ITEM", None)
                or rows_raw.get("Row", None)
                or rows_raw.get("ROW", None)
            )
            if inner is None:
                rows = [rows_raw]
            elif isinstance(inner, dict):
                rows = [inner]
            else:
                rows = list(inner)
        elif isinstance(rows_raw, list):
            rows = rows_raw
        else:
            rows = []

        return {"rows": rows, "count": len(rows), "raw_keys": list(rows[0].keys()) if rows else []}
    except Exception as e:
        return {"error": f"Failed to parse customer orders: {e}", "raw": data}


def parse_order_response(data: dict) -> dict:
    """Extract key fields from the GET_ORDER response into a flat summary."""
    try:
        output = data.get("OutputParameters", data)
        header = output.get("X_HEADER_REC", output.get("HEADER_REC", {}))
        header_val = output.get("X_HEADER_VAL_REC", output.get("HEADER_VAL_REC", {}))
        lines_raw = output.get("X_LINE_TBL", output.get("LINE_TBL", {}))
        lines_val_raw = output.get("X_LINE_VAL_TBL", output.get("LINE_VAL_TBL", {}))
        messages = output.get("X_MSG_DATA", output.get("MSG_DATA", ""))

        # Normalise lines to a list
        if isinstance(lines_raw, dict):
            items = lines_raw.get("X_LINE_TBL_ITEM", lines_raw.get("LINE_TBL_ITEM", []))
            if isinstance(items, dict):
                items = [items]
        elif isinstance(lines_raw, list):
            items = lines_raw
        else:
            items = []

        # Normalise line value records
        if isinstance(lines_val_raw, dict):
            val_items = lines_val_raw.get("X_LINE_VAL_TBL_ITEM", lines_val_raw.get("LINE_VAL_TBL_ITEM", []))
            if isinstance(val_items, dict):
                val_items = [val_items]
        elif isinstance(lines_val_raw, list):
            val_items = lines_val_raw
        else:
            val_items = []

        lines = []
        for idx, item in enumerate(items):
            val = val_items[idx] if idx < len(val_items) else {}
            lines.append({
                "Line": item.get("LINE_NUMBER", ""),
                "Item": item.get("ORDERED_ITEM", ""),
                "Qty": item.get("ORDERED_QUANTITY", ""),
                "UOM": item.get("ORDER_QUANTITY_UOM", ""),
                "Unit Price": item.get("UNIT_SELLING_PRICE", ""),
                "Status": item.get("FLOW_STATUS_CODE", ""),
                # Extended fields for PDF
                "List Price": item.get("UNIT_LIST_PRICE", ""),
                "Line Type ID": item.get("LINE_TYPE_ID", ""),
                "Request Date": item.get("REQUEST_DATE", ""),
                "Promise Date": item.get("PROMISE_DATE", ""),
                "Shipment #": item.get("SHIPMENT_NUMBER", ""),
                "Open Flag": item.get("OPEN_FLAG", ""),
                "Booked Flag": item.get("BOOKED_FLAG", ""),
                "Cancelled Qty": item.get("CANCELLED_QUANTITY", ""),
                "Item Type": item.get("ITEM_TYPE_CODE", ""),
                "Ship To": val.get("SHIP_TO_ORG", ""),
                "Ship From": val.get("SHIP_FROM_ORG", ""),
                "Line Type": val.get("LINE_TYPE", ""),
            })

        # Full header details for PDF
        cname = header_val.get("SOLD_TO_ORG", "") or header.get("SOLD_TO_ORG", header.get("SOLD_TO_ORG_ID", ""))
        header_details = {
            "Order Number": header.get("ORDER_NUMBER", ""),
            "Header ID": header.get("HEADER_ID", ""),
            "Customer PO": header.get("CUST_PO_NUMBER", ""),
            "Flow Status": header.get("FLOW_STATUS_CODE", ""),
            "Order Category": header.get("ORDER_CATEGORY_CODE", ""),
            "Ordered Date": header.get("ORDERED_DATE", ""),
            "Request Date": header.get("REQUEST_DATE", ""),
            "Pricing Date": header.get("PRICING_DATE", ""),
            "Currency": header.get("TRANSACTIONAL_CURR_CODE", ""),
            "Booked Flag": header.get("BOOKED_FLAG", ""),
            "Open Flag": header.get("OPEN_FLAG", ""),
            "Cancelled Flag": header.get("CANCELLED_FLAG", ""),
            "Org ID": header.get("ORG_ID", ""),
            "Order Type": header_val.get("ORDER_TYPE", "") or header.get("ORDER_TYPE_ID", ""),
            "Customer Name": cname,
            "Customer Number": header_val.get("CUSTOMER_NUMBER", ""),
            "Ship To": header_val.get("SHIP_TO_ORG", ""),
            "Ship To Address": header_val.get("SHIP_TO_ADDRESS1", ""),
            "Ship To City": header_val.get("SHIP_TO_CITY", ""),
            "Ship To State": header_val.get("SHIP_TO_STATE", ""),
            "Ship To Country": header_val.get("SHIP_TO_COUNTRY", ""),
            "Bill To": header_val.get("INVOICE_TO_ORG", ""),
            "Bill To Address": header_val.get("INVOICE_TO_ADDRESS1", ""),
            "Salesrep": header_val.get("SALESREP", "") or header.get("SALESREP_ID", ""),
            "Payment Term": header_val.get("PAYMENT_TERM", "") or header.get("PAYMENT_TERM_ID", ""),
            "Price List": header_val.get("PRICE_LIST", "") or header.get("PRICE_LIST_ID", ""),
            "Shipping Method": header_val.get("SHIPPING_METHOD", ""),
            "Freight Terms": header_val.get("FREIGHT_TERMS", ""),
            "FOB Point": header_val.get("FOB_POINT", ""),
        }

        return {
            "order_number": header.get("ORDER_NUMBER", ""),
            "status": header.get("FLOW_STATUS_CODE", ""),
            "ordered_date": header.get("ORDERED_DATE", ""),
            "customer_name": cname,
            "currency": header.get("TRANSACTIONAL_CURR_CODE", ""),
            "lines": lines,
            "messages": messages,
            "header_details": header_details,
        }
    except Exception as e:
        return {"error": f"Failed to parse response: {e}", "raw": data}


# ── In-memory store for last operation data (per-process) ─────
_last_operation = {
    "type": None,          # "order", "create", "customer_orders"
    "data": None,          # parsed result dict
    "raw": None,           # raw API response
    "email_html": "",     # pre-built email HTML
}

# ── Auto-email & Scheduler state ──────────────────────────────
_auto_email = {
    "enabled": False,
    "to": "",
}

_scheduler = {
    "active": False,
    "interval_mins": 60,
    "customer_name": "",
    "to_email": "",
    "last_run": None,
    "last_status": None,
    "runs": 0,
}
_stop_event = threading.Event()
_scheduler_thread = None

# ── Auto-Loop State (iterate all customers, email each) ────────
_auto_loop = {
    "active": False,
    "to_email": "",
    "current_customer": "",
    "total_customers": 0,
    "processed": 0,
    "last_status": None,
    "started_at": None,
}
_stop_loop_event = threading.Event()
_loop_thread = None


def get_all_customer_names() -> list:
    """Fetch distinct customer names from HZ_PARTIES that have sales orders."""
    conn = _get_db_connection()
    cur = conn.cursor()
    sql = """
        SELECT DISTINCT HP.PARTY_NAME AS CUSTOMER_NAME
        FROM HZ_PARTIES HP
        JOIN HZ_CUST_ACCOUNTS HCA ON HCA.PARTY_ID = HP.PARTY_ID
        JOIN OE_ORDER_HEADERS_ALL OOHA ON OOHA.SOLD_TO_ORG_ID = HCA.CUST_ACCOUNT_ID
        WHERE OOHA.ORG_ID = :org_id
        AND OOHA.ORDERED_DATE >= ADD_MONTHS(SYSDATE, -36)
        AND HP.PARTY_NAME IS NOT NULL
        ORDER BY HP.PARTY_NAME
    """
    cur.execute(sql, {"org_id": int(ORG_ID)})
    cols = [c[0] for c in cur.description]
    rows = [dict(zip(cols, r)) for r in cur.fetchall()]
    cur.close()
    conn.close()
    return [r["CUSTOMER_NAME"] for r in rows]


def _auto_loop_worker(to_email: str):
    """Background thread: iterate all customers, send email for each, loop until stopped."""
    _stop_loop_event.clear()
    _auto_loop["active"] = True
    _auto_loop["to_email"] = to_email
    _auto_loop["started_at"] = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    while not _stop_loop_event.is_set():
        try:
            customers = get_all_customer_names()
            _auto_loop["total_customers"] = len(customers)
            _auto_loop["processed"] = 0
            print(f"Auto-loop: Starting cycle with {len(customers)} customers -> {to_email}")

            for cust_name in customers:
                if _stop_loop_event.is_set():
                    break

                _auto_loop["current_customer"] = cust_name
                try:
                    raw = get_customer_orders(cust_name)
                    parsed = parse_customer_orders(raw)
                    parsed["customer_name"] = cust_name
                    rows = parsed.get("rows", [])

                    if rows:
                        # Store in _last_operation so _generate_excel() picks it up
                        _last_operation["type"] = "customer_orders"
                        _last_operation["data"] = parsed
                        _last_operation["raw"] = raw

                        html_payload = _generate_customer_orders_email_html(parsed)
                        _last_operation["email_html"] = html_payload
                        subject = f"Sales Order Report: {cust_name}"
                        body_text = f"Automated Sales Order Report for customer: {cust_name}"
                        success, msg = send_erp_email(to_email, subject, body_text, html_payload)
                        _auto_loop["last_status"] = f"OK {cust_name}: {msg}" if success else f"FAIL {cust_name}: {msg}"
                        print(f"  {'OK' if success else 'FAIL'} {cust_name} ({len(rows)} rows) -> {msg}")
                    else:
                        _auto_loop["last_status"] = f"SKIP {cust_name}: No orders"
                        print(f"  SKIP {cust_name}: No orders")
                except Exception as e:
                    _auto_loop["last_status"] = f"ERR {cust_name}: {e}"
                    print(f"  ERR {cust_name}: {e}")

                _auto_loop["processed"] += 1

                if not _stop_loop_event.is_set():
                    _stop_loop_event.wait(timeout=5)

            if not _stop_loop_event.is_set():
                _auto_loop["last_status"] = f"Cycle complete ({len(customers)} customers). Waiting 60s..."
                print(f"Auto-loop: Cycle complete. Waiting 60s...")
                _stop_loop_event.wait(timeout=60)

        except Exception as e:
            _auto_loop["last_status"] = f"Loop error: {e}"
            print(f"Auto-loop error: {e}")
            if not _stop_loop_event.is_set():
                _stop_loop_event.wait(timeout=30)

    _auto_loop["active"] = False
    _auto_loop["current_customer"] = ""
    print("Auto-loop: Stopped.")


def _start_auto_loop(to_email: str):
    """Start the auto email loop in a background thread."""
    global _loop_thread
    _stop_auto_loop()
    _loop_thread = threading.Thread(target=_auto_loop_worker, args=(to_email,), daemon=True)
    _loop_thread.start()


def _stop_auto_loop():
    """Stop the auto email loop."""
    global _loop_thread
    _stop_loop_event.set()
    _auto_loop["active"] = False
    if _loop_thread and _loop_thread.is_alive():
        _loop_thread.join(timeout=3)
    _loop_thread = None
    _stop_loop_event.clear()


def _scheduled_worker(interval_mins: int, customer_name: str, to_email: str):
    """Background thread: fetch order data and email on a recurring interval."""
    _stop_event.clear()
    while not _stop_event.is_set():
        stopped = _stop_event.wait(timeout=interval_mins * 60)
        if stopped:
            break
        try:
            target = customer_name.strip() if customer_name else ""
            if target:
                raw = get_customer_orders(target)
                parsed = parse_customer_orders(raw)
                parsed["customer_name"] = target
                html_payload = _generate_customer_orders_email_html(parsed)
                subject = f"Scheduled Report: {target}"
            else:
                html_payload = _last_operation.get("email_html", "")
                subject = "Scheduled Report: Latest Order Data"

            if html_payload:
                body_text = "Your automated scheduled report is attached below."
                success, msg = send_erp_email(to_email, subject, body_text, html_payload)
                _scheduler["last_run"] = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
                _scheduler["runs"] += 1
                _scheduler["last_status"] = "sent" if success else f"failed: {msg}"
            else:
                _scheduler["last_run"] = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
                _scheduler["last_status"] = "skipped: no data"
        except Exception as e:
            _scheduler["last_run"] = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
            _scheduler["last_status"] = f"error: {e}"


def _start_scheduler(interval_mins: int, customer_name: str, to_email: str):
    """Start the scheduler background thread."""
    global _scheduler_thread
    _stop_scheduler()
    _scheduler["active"] = True
    _scheduler["interval_mins"] = interval_mins
    _scheduler["customer_name"] = customer_name
    _scheduler["to_email"] = to_email
    _scheduler["runs"] = 0
    _scheduler["last_run"] = None
    _scheduler["last_status"] = None
    _scheduler_thread = threading.Thread(
        target=_scheduled_worker,
        args=(interval_mins, customer_name, to_email),
        daemon=True,
    )
    _scheduler_thread.start()


def _stop_scheduler():
    """Stop the scheduler background thread."""
    global _scheduler_thread
    _stop_event.set()
    _scheduler["active"] = False
    if _scheduler_thread and _scheduler_thread.is_alive():
        _scheduler_thread.join(timeout=2)
    _scheduler_thread = None
    _stop_event.clear()


def trigger_auto_email(context: str):
    """Send auto-email if enabled, using the last operation's HTML."""
    if not _auto_email["enabled"] or not _auto_email["to"]:
        return
    html_payload = _last_operation.get("email_html", "")
    if not html_payload:
        return
    to_email = _auto_email["to"]
    subject = f"Automated ERP Update — {context}"
    body_text = (
        "Hi,\n\nPlease find the automated report attached below.\n\n"
        "This email was sent automatically after the operation completed.\n\n"
        "Regards,\nSales Order Assistant"
    )
    threading.Thread(
        target=send_erp_email,
        args=(to_email, subject, body_text, html_payload),
        daemon=True,
    ).start()


# ── Email Template (matches Stream2.py layout) ───────────────
def _build_email_layout(cname: str, report_type: str, ai_summary: str,
                        table_html: str, alerts: dict, faqs: list) -> str:
    report_date = datetime.datetime.now().strftime("%Y-%m-%d")
    alerts_html = "<ul style='padding-left: 20px; font-size: 14px; margin-bottom: 0;'>"
    for k, v in alerts.items():
        alerts_html += f"<li style='margin-bottom: 5px;'>{k}: {v}</li>"
    alerts_html += "</ul>"

    faqs_html = (
        "<table style='border-collapse: collapse; width: 100%; border: none; font-size: 14px;'>"
        "<tr><th style='text-align: left; padding: 8px; border-bottom: 1px solid #ccc;'>Question</th>"
        "<th style='text-align: left; padding: 8px; border-bottom: 1px solid #ccc;'>AI Answer</th></tr>"
    )
    for q, a in faqs:
        faqs_html += f"<tr><td style='padding: 8px;'>{q}</td><td style='padding: 8px;'>{a}</td></tr>"
    faqs_html += "</table>"

    return f"""
    <div style="font-family: Arial, sans-serif; max-width: 800px; margin: auto; color: #000; padding: 10px;">

        <h3 style="color: #000; font-size: 16px; margin-top: 0;">Header Details</h3>
        <ul style="list-style-type: disc; padding-left: 20px; font-size: 14px; margin-bottom: 20px;">
            <li style="margin-bottom: 5px;">Customer Name: {cname}</li>
            <li style="margin-bottom: 5px;">Contact Name: Account Primary</li>
            <li style="margin-bottom: 5px;">Date of Report: {report_date}</li>
            <li style="margin-bottom: 5px;">Statement Period: {report_type}</li>
        </ul>
        <hr style="border: 0; border-top: 1px solid #aaa; margin: 20px 0;">

        <h3 style="color: #000; font-size: 16px;">Executive Summary</h3>
        <p>{ai_summary}</p>
        <hr style="border: 0; border-top: 1px solid #aaa; margin: 20px 0;">

        <h3 style="color: #000; font-size: 16px;">Order Detail Table</h3>
        {table_html}
        <hr style="border: 0; border-top: 1px solid #aaa; margin: 20px 0;">

        <h3 style="color: #000; font-size: 16px;">Exceptions & Alerts</h3>
        {alerts_html}
        <hr style="border: 0; border-top: 1px solid #aaa; margin: 20px 0;">

        <h3 style="color: #000; font-size: 16px;">FAQ'S</h3>
        {faqs_html}
        <hr style="border: 0; border-top: 1px solid #aaa; margin: 20px 0;">

        <ul style="list-style-type: disc; padding-left: 20px; font-size: 14px; margin-bottom: 30px;">
            <li style="margin-bottom: 5px;">Support contact: erp.support@centroid.com</li>
            <li style="margin-bottom: 5px;">Disclaimer: This notification was generated automatically from the ERP environment</li>
            <li style="margin-bottom: 5px;">Do-not-reply notice: Please do not reply directly to this email.</li>
        </ul>
    </div>
    """


def _generate_order_email_html(parsed: dict) -> str:
    """Build email HTML from a GET_ORDER parsed result."""
    order_num = parsed.get("order_number", "N/A")
    cname = parsed.get("customer_name", "Unknown Customer")
    status = parsed.get("status", "N/A")
    th = "border: none; padding: 8px; text-align: left; font-weight: bold;"
    td = "border: none; padding: 8px;"
    table = f"<table style='border-collapse: collapse; width: 100%; font-size: 14px;'><tr>"
    for h in ["Order #", "Line Item", "Qty", "Status", "Ship Date", "Tracking", "Invoice Balance"]:
        table += f"<th style='{th}'>{h}</th>"
    table += "</tr>"
    for ln in parsed.get("lines", []):
        table += "<tr>"
        table += f"<td style='{td}'>{order_num}</td>"
        table += f"<td style='{td}'>{ln.get('Item', 'N/A')}</td>"
        table += f"<td style='{td}'>{ln.get('Qty', 'N/A')}</td>"
        table += f"<td style='{td}'>{ln.get('Status', 'N/A')}</td>"
        table += f"<td style='{td}'>TBD</td>"
        table += f"<td style='{td}'>Pending</td>"
        table += f"<td style='{td}'>N/A</td>"
        table += "</tr>"
    table += "</table>"
    has_holds = status in ("ENTERED",) or "HOLD" in str(status).upper()
    alerts = {
        "Backorders": "None detected for this specific transaction",
        "Holds": "1" if has_holds else "0",
        "Partial Shipments": "N/A",
        "Payment Holds": "Check customer balance",
    }
    faqs = [
        (f"Why is Order {order_num} currently {status}?",
         f"The order is in {status} status waiting for the next step in the fulfillment cycle."),
        ("Can I expect partial shipment?",
         "Depends on line-level availability. Check shipping notices for updates."),
    ]
    summary = f"Order #{order_num} for customer {cname} is currently in {status} status."
    return _build_email_layout(cname, f"Transaction: Order #{order_num}", summary, table, alerts, faqs)


def _generate_create_email_html(parsed: dict) -> str:
    """Build email HTML from a PROCESS_ORDER (create) result."""
    order_num = parsed.get("order_number", "N/A")
    status = parsed.get("flow_status", "N/A")
    cust_po = parsed.get("cust_po_number", "N/A")
    th = "border: none; padding: 8px; text-align: left; font-weight: bold;"
    td = "border: none; padding: 8px;"
    table = f"<table style='border-collapse: collapse; width: 100%; font-size: 14px;'><tr>"
    for h in ["Order #", "Customer PO", "Status", "Return Status"]:
        table += f"<th style='{th}'>{h}</th>"
    table += "</tr>"
    table += f"<tr><td style='{td}'>{order_num}</td><td style='{td}'>{cust_po}</td>"
    table += f"<td style='{td}'>{status}</td><td style='{td}'>{parsed.get('return_status', '')}</td></tr>"
    table += "</table>"
    alerts = {"Backorders": "N/A", "Holds": "0", "Partial Shipments": "N/A", "Payment Holds": "N/A"}
    faqs = [("Was the order booked?", f"Order #{order_num} has status: {status}.")]
    summary = f"New order #{order_num} was created with PO {cust_po}. Current status: {status}."
    return _build_email_layout("N/A", f"New Order Created: #{order_num}", summary, table, alerts, faqs)


def _generate_customer_orders_email_html(parsed: dict) -> str:
    """Build email HTML from a CustSODtls result."""
    cname = parsed.get("customer_name", "Unknown Customer")
    rows = parsed.get("rows", [])
    keys = parsed.get("raw_keys", [])
    th = "border: none; padding: 8px; text-align: left; font-weight: bold;"
    td = "border: none; padding: 8px;"
    # Use up to 7 columns or all if fewer
    display_keys = keys[:7] if len(keys) > 7 else keys
    table = f"<table style='border-collapse: collapse; width: 100%; font-size: 14px;'><tr>"
    for h in display_keys:
        table += f"<th style='{th}'>{h}</th>"
    table += "</tr>"
    for r in rows[:50]:
        table += "<tr>"
        for k in display_keys:
            table += f"<td style='{td}'>{r.get(k, 'N/A')}</td>"
        table += "</tr>"
    table += "</table>"
    alerts = {"Backorders": "0", "Holds": "0", "Partial Shipments": "Check line statuses", "Payment Holds": "0"}
    faqs = [("Are there any delayed orders?", "Check the table for delayed order flags.")]
    summary = f"{len(rows)} order record(s) found for customer {cname}."
    return _build_email_layout(cname, "All Recent Customer Orders", summary, table, alerts, faqs)


def _generate_orders_advanced_email_html(data: dict, range_type=None, year=None, status=None) -> str:
    """Build email HTML from an orders_advanced result."""
    orders = data.get("orders", [])
    total = data.get("total_count", 0)
    label_parts = []
    if range_type:
        label_parts.append(range_type.replace("_", " ").title())
    if year:
        label_parts.append(str(year))
    if status:
        label_parts.append(status)
    filter_label = ", ".join(label_parts) or "All"

    th = "border: none; padding: 8px; text-align: left; font-weight: bold;"
    td = "border: none; padding: 8px;"
    table = "<table style='border-collapse: collapse; width: 100%; font-size: 14px;'><tr>"
    for h in ["Order #", "Header Status", "Line Status", "Creation Date"]:
        table += f"<th style='{th}'>{h}</th>"
    table += "</tr>"
    for o in orders[:50]:
        table += f"<tr><td style='{td}'>{o.get('order_number','')}</td>"
        table += f"<td style='{td}'>{o.get('header_status','')}</td>"
        table += f"<td style='{td}'>{o.get('line_status','')}</td>"
        table += f"<td style='{td}'>{o.get('creation_date','')}</td></tr>"
    table += "</table>"
    alerts = {"Total Orders": str(total), "Filter": filter_label}
    faqs = [("What filter was applied?", f"Filter: {filter_label}")]
    summary_text = f"{total} order(s) found with filter: {filter_label}."
    return _build_email_layout("N/A", f"Orders Report: {filter_label}", summary_text, table, alerts, faqs)


# ── Excel Generation ────────────────────────────────
import io


def _generate_excel() -> bytes:
    """Generate an Excel workbook from the last operation data."""
    if not OPENPYXL_AVAILABLE:
        return b""

    op_type = _last_operation.get("type")
    parsed = _last_operation.get("data")
    if not op_type or not parsed:
        return b""

    wb = Workbook()
    header_font = Font(bold=True, color="FFFFFF", size=10)
    header_fill = PatternFill(start_color="003366", end_color="003366", fill_type="solid")
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    def style_header(ws, row_num, col_count):
        for col in range(1, col_count + 1):
            cell = ws.cell(row=row_num, column=col)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thin_border

    def style_data(ws, start_row, end_row, col_count):
        for row in range(start_row, end_row + 1):
            for col in range(1, col_count + 1):
                cell = ws.cell(row=row, column=col)
                cell.border = thin_border
                cell.alignment = Alignment(vertical="center")

    def auto_width(ws):
        for col in ws.columns:
            max_len = 0
            col_letter = col[0].column_letter
            for cell in col:
                if cell.value:
                    max_len = max(max_len, len(str(cell.value)))
            ws.column_dimensions[col_letter].width = min(max_len + 3, 40)

    if op_type == "order":
        # Sheet 1: Header Details
        ws = wb.active
        ws.title = "Header Details"
        hdr = parsed.get("header_details", {})
        ws.append(["Field", "Value"])
        style_header(ws, 1, 2)
        row_num = 2
        for key, val in hdr.items():
            val_str = str(val) if val not in (None, "", "None") else ""
            if val_str:
                ws.append([key, val_str])
                row_num += 1
        style_data(ws, 2, row_num, 2)
        auto_width(ws)

        # Sheet 2: Line Details
        lines = parsed.get("lines", [])
        if lines:
            ws2 = wb.create_sheet("Line Details")
            cols = list(lines[0].keys())
            ws2.append(cols)
            style_header(ws2, 1, len(cols))
            for i, line in enumerate(lines):
                ws2.append([line.get(c, "") for c in cols])
            style_data(ws2, 2, len(lines) + 1, len(cols))
            auto_width(ws2)

    elif op_type == "create":
        ws = wb.active
        ws.title = "Order Details"
        ws.append(["Field", "Value"])
        style_header(ws, 1, 2)
        row_num = 2
        for key, val in parsed.items():
            if key == "success":
                continue
            val_str = str(val) if val not in (None, "", "None") else ""
            if val_str:
                ws.append([key, val_str])
                row_num += 1
        style_data(ws, 2, row_num, 2)
        auto_width(ws)

    elif op_type == "customer_orders":
        ws = wb.active
        ws.title = "Customer Orders"
        rows = parsed.get("rows", [])
        keys = parsed.get("raw_keys", [])
        if not keys and rows:
            keys = list(rows[0].keys())
        if keys:
            ws.append(keys)
            style_header(ws, 1, len(keys))
            for i, row in enumerate(rows):
                ws.append([row.get(k, "") for k in keys])
            style_data(ws, 2, len(rows) + 1, len(keys))
            auto_width(ws)

    elif op_type == "orders_advanced":
        ws = wb.active
        ws.title = "Orders Report"
        orders = parsed.get("orders", [])
        cols = ["order_number", "header_status", "line_status", "creation_date"]
        ws.append(["Order #", "Header Status", "Line Status", "Creation Date"])
        style_header(ws, 1, len(cols))
        for o in orders:
            ws.append([o.get(c, "") for c in cols])
        style_data(ws, 2, len(orders) + 1, len(cols))
        auto_width(ws)

    else:
        return b""

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _get_excel_filename() -> str:
    """Generate a descriptive Excel filename."""
    op = _last_operation.get("type") or "report"
    data = _last_operation.get("data") or {}
    ts = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
    if op == "order":
        return f"Order_{data.get('order_number', 'N-A')}_{ts}.xlsx"
    elif op == "customer_orders":
        cname = data.get("customer_name", "unknown").replace(" ", "_")
        return f"CustOrders_{cname}_{ts}.xlsx"
    elif op == "create":
        return f"NewOrder_{data.get('order_number', 'N-A')}_{ts}.xlsx"
    elif op == "orders_advanced":
        return f"OrdersReport_{ts}.xlsx"
    return f"ERP_Report_{ts}.xlsx"


def send_erp_email(to_email: str, subject: str, body_text: str, html_table: str) -> tuple:
    """Send an email with the formatted HTML report and attached Excel file."""
    from email.mime.application import MIMEApplication
    sender = "erp.assistant@yourdomain.com"
    msg = MIMEMultipart("mixed")
    msg["Subject"] = subject
    msg["From"] = sender
    msg["To"] = to_email

    html_content = f"""
    <html>
    <head></head>
    <body style="font-family: Arial, sans-serif; background-color: #f4f4f4; padding: 20px;">
        <p style="color: #333;">{body_text.replace(chr(10), '<br>')}</p>
        <br>
        {html_table}
    </body>
    </html>
    """

    # Attach text and HTML as alternatives inside a sub-part
    alt_part = MIMEMultipart("alternative")
    alt_part.attach(MIMEText(body_text, "plain"))
    alt_part.attach(MIMEText(html_content, "html"))
    msg.attach(alt_part)

    # Attach Excel file
    excel_bytes = _generate_excel()
    if excel_bytes:
        excel_name = _get_excel_filename()
        excel_part = MIMEApplication(excel_bytes, _subtype="vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        excel_part.add_header("Content-Disposition", "attachment", filename=excel_name)
        msg.attach(excel_part)

    try:
        server = smtplib.SMTP("smtp.gmail.com", 587)
        server.starttls()
        server.login("pavanmaccha042@gmail.com", "ltvxbwhkzkrptueg")
        server.sendmail(sender, to_email, msg.as_string())
        server.quit()
        return True, "Email sent successfully!"
    except Exception as e:
        return False, str(e)


# ── Flask App ─────────────────────────────────────
app = Flask(__name__, static_folder="static", static_url_path="/static")


@app.route("/")
def index():
    return send_from_directory("static", "index5.html")


@app.route("/api/chat", methods=["POST"])
def chat():
    """Handle a chat message from the user."""
    body = request.get_json(force=True)
    user_msg = (body.get("message") or "").strip()

    if not user_msg:
        return jsonify({"reply": "Please type a message.", "type": "text"})

    # Normalize voice input (convert spoken numbers to digits, etc.)
    normalized = normalize_voice_text(user_msg)
    print(f"\U0001f4ac  Input: '{user_msg}' -> Normalized: '{normalized}'")

    lower = normalized.lower()

    # Check for greetings
    greetings = ["hi", "hello", "hey", "good morning", "good afternoon", "good evening"]
    if lower in greetings:
        return jsonify({
            "reply": "Hello! I'm the Sales Order Assistant.\n\nYou can:\n\u2022 **Get order 67965** \u2014 Look up a sales order\n\u2022 **Customer orders for Vision Operations** \u2014 Orders by customer\n\u2022 **Create order** \u2014 Create a new sales order\n\u2022 **Orders this month** / **Orders in 2024** / **Booked orders** \u2014 Filter orders\n\u2022 **Add line to order 67965** \u2014 Add item to existing order\n\u2022 **Send email to user@domain.com** \u2014 Email last operation results\n\u2022 Or ask me any general question!",
            "type": "text",
        })

    # Check for help
    if lower in ("help", "?", "commands"):
        return jsonify({
            "reply": "Here's what I can do:\n\n\u2022 **Get order [number]** \u2014 Retrieve details for a sales order\n\u2022 **Customer orders for [name]** \u2014 Get all orders for a customer\n\u2022 **Create order** \u2014 Create a new sales order\n\u2022 **Orders this month** / **Orders in 2024** / **Booked orders** \u2014 Filter by date/year/status\n\u2022 **Add line to order [number]** \u2014 Add a line item to an existing order\n\u2022 **Send email to [address]** \u2014 Email the last operation results\n\u2022 Or ask me any general question!",
            "type": "text",
        })

    import re

    # ── Quick regex checks for email (avoid AI call) ──
    email_match = re.search(
        r"(?:send|email|mail)\s+(?:email\s+)?(?:to\s+)?([\w.+-]+@[\w.-]+\.\w+)",
        normalized, re.IGNORECASE,
    )
    if email_match:
        to_addr = email_match.group(1)
        return jsonify({
            "reply": {"to_email": to_addr, "has_data": _last_operation["type"] is not None},
            "type": "email_form",
        })

    if re.search(r"\b(send|email|mail)\b", lower) and not re.search(r"\b(create|new|place|book|add|get|show|order|customer)\b", lower):
        return jsonify({
            "reply": {"to_email": "", "has_data": _last_operation["type"] is not None},
            "type": "email_form",
        })

    # ── Use AI to extract intent ──
    intent = ai_extract_intent(normalized)
    action = intent.get("action", "unknown")
    print(f"\U0001f9e0  AI Intent: {json.dumps(intent)}")

    # ── Action: create order ──
    if action == "create":
        prefilled = {}
        # Extract any inline create-order fields from the AI intent
        cpo = intent.get("cust_po_number")
        item_name = intent.get("ordered_item")
        qty = intent.get("ordered_quantity")
        cust_name = intent.get("customer_name")

        if cpo:
            prefilled["cust_po_number"] = str(cpo)

        if qty:
            try:
                prefilled["ordered_quantity"] = float(qty)
            except (ValueError, TypeError):
                pass

        if cust_name:
            prefilled["customer_name"] = str(cust_name)

        # Validate item by looking it up in the database
        if item_name:
            prefilled["ordered_item_input"] = str(item_name)
            try:
                conn = _get_db_connection()
                cur = conn.cursor()
                # Try exact match first
                sql_exact = """
                    SELECT INVENTORY_ITEM_ID, SEGMENT1, DESCRIPTION
                      FROM MTL_SYSTEM_ITEMS_B
                     WHERE UPPER(SEGMENT1) = :item_name AND ORGANIZATION_ID = :org_id
                     AND ROWNUM = 1
                """
                cur.execute(sql_exact, item_name=str(item_name).upper(), org_id=int(ORG_ID))
                row = cur.fetchone()
                if row:
                    prefilled["inventory_item_id"] = int(row[0])
                    prefilled["ordered_item"] = row[1]
                    prefilled["item_description"] = row[2] or ""
                    prefilled["item_validated"] = True
                else:
                    prefilled["item_validated"] = False
                cur.close()
                conn.close()
            except Exception as e:
                print(f"Item validation error: {e}")
                prefilled["item_validated"] = False

        # Validate customer by name
        if cust_name:
            try:
                conn = _get_db_connection()
                cur = conn.cursor()
                sql_cust = """
                    SELECT hca.CUST_ACCOUNT_ID, hp.PARTY_NAME, hca.ACCOUNT_NUMBER
                      FROM HZ_CUST_ACCOUNTS hca
                      JOIN HZ_PARTIES hp ON hp.PARTY_ID = hca.PARTY_ID
                     WHERE UPPER(hp.PARTY_NAME) LIKE :cust_name AND hca.STATUS = 'A'
                     AND ROWNUM <= 5
                     ORDER BY hp.PARTY_NAME
                """
                cur.execute(sql_cust, cust_name=f"%{cust_name.upper()}%")
                cust_rows = cur.fetchall()
                if cust_rows and len(cust_rows) == 1:
                    prefilled["cust_account_id"] = int(cust_rows[0][0])
                    prefilled["customer_name_validated"] = cust_rows[0][1]
                    prefilled["customer_validated"] = True
                else:
                    prefilled["customer_validated"] = False
                cur.close()
                conn.close()
            except Exception as e:
                print(f"Customer validation error: {e}")
                prefilled["customer_validated"] = False

        return jsonify({"reply": "create_form", "type": "create_form", "prefilled": prefilled})

    # ── Action: add line to order ──
    if action == "add_line":
        order_num = intent.get("order_number")
        if order_num:
            # Fetch the order to get header_id
            try:
                raw = get_order(int(order_num))
                parsed = parse_order_response(raw)
                header_id = ""
                try:
                    out = raw.get("PROCESS_ORDER_Output", {}).get("OutputParameters", {})
                    hdr = out.get("X_HEADER_REC", {})
                    header_id = hdr.get("HEADER_ID") or hdr.get("header_id") or ""
                    if not header_id:
                        out2 = raw.get("OutputParameters", {})
                        hdr2 = out2.get("X_HEADER_REC", {})
                        header_id = hdr2.get("HEADER_ID") or ""
                except Exception:
                    pass
                return jsonify({
                    "reply": {"order_number": int(order_num), "header_id": header_id},
                    "type": "add_line_form",
                })
            except Exception as e:
                return jsonify({"reply": f"Could not fetch order {order_num}: {e}", "type": "error"})
        else:
            return jsonify({
                "reply": {"order_number": None, "header_id": None},
                "type": "add_line_form",
            })

    # ── Action: get_range (orders by date/year/status) ──
    if action == "get_range":
        range_type = intent.get("date_range")
        year = intent.get("year")
        status = intent.get("status")
        try:
            data = get_orders_advanced(range_type=range_type, year=year, status=status)
            _last_operation["type"] = "orders_advanced"
            _last_operation["data"] = data
            _last_operation["raw"] = data
            label_parts = []
            if range_type:
                label_parts.append(range_type.replace("_", " "))
            if year:
                label_parts.append(str(year))
            if status:
                label_parts.append(status)
            label = ", ".join(label_parts) or "all"
            matched_on = data.get("matched_on")
            date_filter_dropped = data.get("date_filter_dropped", False)
            if matched_on and status:
                if matched_on == "header_status":
                    label += " (matched on Header Status)"
                elif matched_on == "line_status":
                    label += " (matched on Line Status)"
                elif matched_on == "none":
                    label += " (no match in Header or Line Status)"
            if date_filter_dropped:
                label += " ⚠️ No orders for requested date range — showing all dates"
            _last_operation["email_html"] = _generate_orders_advanced_email_html(data, range_type, year, status)
            summary = generate_ai_summary("orders_advanced", data)
            _last_operation["ai_summary"] = summary
            return jsonify({"reply": data, "type": "orders_advanced", "ai_summary": summary, "filter_label": label, "matched_on": matched_on})
        except Exception as e:
            return jsonify({"reply": f"Error querying orders: {e}", "type": "error"})

    # ── Action: get (specific order or customer orders) ──
    if action == "get":
        customer_name = intent.get("customer_name")
        order_number = intent.get("order_number")

        if customer_name and not order_number:
            try:
                raw = get_customer_orders(customer_name)
                parsed = parse_customer_orders(raw)
                parsed["customer_name"] = customer_name
                _last_operation["type"] = "customer_orders"
                _last_operation["data"] = parsed
                _last_operation["raw"] = raw
                _last_operation["email_html"] = _generate_customer_orders_email_html(parsed)
                trigger_auto_email(f"Customer: {customer_name}")
                summary = generate_ai_summary("customer_orders", parsed)
                _last_operation["ai_summary"] = summary
                return jsonify({"reply": parsed, "type": "customer_orders", "ai_summary": summary})
            except requests.exceptions.ConnectionError:
                return jsonify({"reply": "Could not connect to the Oracle EBS server. Please check network/VPN.", "type": "error"})
            except Exception as e:
                return jsonify({"reply": f"Something went wrong: {e}", "type": "error"})

        if order_number:
            try:
                raw = get_order(int(order_number))
                parsed = parse_order_response(raw)
                _last_operation["type"] = "order"
                _last_operation["data"] = parsed
                _last_operation["raw"] = raw
                _last_operation["email_html"] = _generate_order_email_html(parsed)
                trigger_auto_email(f"Order #{order_number}")
                summary = generate_ai_summary("order", parsed)
                _last_operation["ai_summary"] = summary
                return jsonify({"reply": parsed, "type": "order", "ai_summary": summary})
            except requests.exceptions.ConnectionError:
                return jsonify({"reply": "Could not connect to the Oracle EBS server. Please check network/VPN.", "type": "error"})
            except Exception as e:
                return jsonify({"reply": f"Something went wrong: {e}", "type": "error"})

        return jsonify({"reply": "Which order number would you like to look up? Please say or type the number.", "type": "text"})

    # ── Action: email ──
    if action == "email":
        email_addr = intent.get("email_address", "")
        return jsonify({
            "reply": {"to_email": email_addr or "", "has_data": _last_operation["type"] is not None},
            "type": "email_form",
        })

    # ── Action: general (AI answers general questions) ──
    if action == "general":
        answer = intent.get("general_answer", "")
        if answer:
            return jsonify({"reply": answer, "type": "text"})

    # ── Fallback: regex-based order number extraction ──
    order_match = re.search(r"\b(\d{4,})\b", normalized)
    if order_match:
        order_num = int(order_match.group(1))
        try:
            raw = get_order(order_num)
            parsed = parse_order_response(raw)
            _last_operation["type"] = "order"
            _last_operation["data"] = parsed
            _last_operation["raw"] = raw
            _last_operation["email_html"] = _generate_order_email_html(parsed)
            trigger_auto_email(f"Order #{order_num}")
            summary = generate_ai_summary("order", parsed)
            _last_operation["ai_summary"] = summary
            return jsonify({"reply": parsed, "type": "order", "ai_summary": summary})
        except Exception as e:
            return jsonify({"reply": f"Something went wrong: {e}", "type": "error"})

    # ── Fallback: customer order regex ──
    cust_match = re.search(
        r"(?:customer\s+(?:orders?|details?)\s+(?:(?:for|of)\s+)?(.+)"
        r"|orders?\s+(?:for|of|by)\s+(?:customer\s+)?(.+)"
        r"|(?:show|get|find)\s+customer\s+orders?\s*(?:(?:for|of|by)\s+)?(.+))",
        normalized, re.IGNORECASE,
    )
    if cust_match:
        cust_name = (cust_match.group(1) or cust_match.group(2) or cust_match.group(3) or "").strip().strip("'\"")
        cust_name = re.sub(r"\s+(please|thanks|thank you)$", "", cust_name, flags=re.IGNORECASE).strip()
        if cust_name:
            try:
                raw = get_customer_orders(cust_name)
                parsed = parse_customer_orders(raw)
                parsed["customer_name"] = cust_name
                _last_operation["type"] = "customer_orders"
                _last_operation["data"] = parsed
                _last_operation["raw"] = raw
                _last_operation["email_html"] = _generate_customer_orders_email_html(parsed)
                summary = generate_ai_summary("customer_orders", parsed)
                _last_operation["ai_summary"] = summary
                return jsonify({"reply": parsed, "type": "customer_orders", "ai_summary": summary})
            except Exception as e:
                return jsonify({"reply": f"Something went wrong: {e}", "type": "error"})

    return jsonify({
        "reply": "I'm not sure what you need. Here are some things I can help with:\n\n\u2022 **Get order 67965** \u2014 Look up a sales order\n\u2022 **Customer orders for Vision Operations** \u2014 Orders by customer\n\u2022 **Create order** \u2014 Create a new sales order\n\u2022 **Orders this month** / **Booked orders** \u2014 Filter orders\n\u2022 **Add line to order 67965** \u2014 Add line item\n\u2022 **Send email** \u2014 Email results\n\nOr ask me any general question!",
        "type": "text",
    })


# ── Lookup API endpoints for create-order wizard ──

@app.route("/api/lookup-items", methods=["GET"])
def api_lookup_items():
    """Search inventory items by prefix."""
    search = request.args.get("search", "").strip()
    offset = int(request.args.get("offset", 0))
    if not search:
        return jsonify({"items": [], "has_more": False})
    try:
        conn = _get_db_connection()
        cur = conn.cursor()
        sql = """
            SELECT * FROM (
                SELECT a.*, ROWNUM rnum FROM (
                    SELECT DISTINCT INVENTORY_ITEM_ID, SEGMENT1, DESCRIPTION, ORGANIZATION_ID
                      FROM MTL_SYSTEM_ITEMS_B
                     WHERE SEGMENT1 LIKE :ordered_item AND ORGANIZATION_ID = :org_id
                     ORDER BY INVENTORY_ITEM_ID
                ) a WHERE ROWNUM <= :upper_bound
            ) WHERE rnum > :offset
        """
        cur.execute(sql, ordered_item=f"{search.upper()}%", org_id=int(ORG_ID),
                    upper_bound=offset + MAX_CHOICES, offset=offset)
        rows = cur.fetchall()
        cur.close()
        conn.close()
        items = [{"inventory_item_id": int(r[0]), "segment1": r[1], "description": r[2] or ""} for r in rows]
        return jsonify({"items": items, "has_more": len(items) == MAX_CHOICES})
    except Exception as e:
        return jsonify({"items": [], "has_more": False, "error": str(e)})


@app.route("/api/lookup-customers", methods=["GET"])
def api_lookup_customers():
    """Search customers by name or by inventory item."""
    search = request.args.get("search", "").strip()
    item_id = request.args.get("item_id", "").strip()
    offset = int(request.args.get("offset", 0))
    try:
        conn = _get_db_connection()
        cur = conn.cursor()
        if search:
            sql = """
                SELECT * FROM (
                    SELECT a.*, ROWNUM rnum FROM (
                        SELECT DISTINCT hca.CUST_ACCOUNT_ID, hp.PARTY_NAME AS customer_name, hca.ACCOUNT_NUMBER
                          FROM HZ_CUST_ACCOUNTS hca
                          JOIN HZ_PARTIES hp ON hp.PARTY_ID = hca.PARTY_ID
                         WHERE UPPER(hp.PARTY_NAME) LIKE :cust_name AND hca.STATUS = 'A'
                         ORDER BY hp.PARTY_NAME
                    ) a WHERE ROWNUM <= :upper_bound
                ) WHERE rnum > :offset
            """
            cur.execute(sql, cust_name=f"%{search.upper()}%",
                        upper_bound=offset + MAX_CHOICES, offset=offset)
        elif item_id:
            sql = """
                SELECT * FROM (
                    SELECT a.*, ROWNUM rnum FROM (
                        SELECT DISTINCT hca.CUST_ACCOUNT_ID, hp.PARTY_NAME AS customer_name, hca.ACCOUNT_NUMBER
                          FROM OE_ORDER_LINES_ALL ool
                          JOIN OE_ORDER_HEADERS_ALL ooh ON ooh.HEADER_ID = ool.HEADER_ID
                          JOIN HZ_CUST_ACCOUNTS hca ON hca.CUST_ACCOUNT_ID = ooh.SOLD_TO_ORG_ID
                          JOIN HZ_PARTIES hp ON hp.PARTY_ID = hca.PARTY_ID
                         WHERE ool.INVENTORY_ITEM_ID = :item_id AND ool.ORG_ID = :org_id AND hca.STATUS = 'A'
                         ORDER BY hp.PARTY_NAME
                    ) a WHERE ROWNUM <= :upper_bound
                ) WHERE rnum > :offset
            """
            cur.execute(sql, item_id=int(item_id), org_id=int(ORG_ID),
                        upper_bound=offset + MAX_CHOICES, offset=offset)
        else:
            cur.close()
            conn.close()
            return jsonify({"customers": [], "has_more": False})
        rows = cur.fetchall()
        cur.close()
        conn.close()
        customers = [{"cust_account_id": int(r[0]), "customer_name": r[1] or "", "account_number": r[2] or ""} for r in rows]
        return jsonify({"customers": customers, "has_more": len(customers) == MAX_CHOICES})
    except Exception as e:
        return jsonify({"customers": [], "has_more": False, "error": str(e)})


@app.route("/api/lookup-customer-sites", methods=["GET"])
def api_lookup_customer_sites():
    """Get ship-to sites for a customer account."""
    cust_account_id = request.args.get("cust_account_id", "").strip()
    if not cust_account_id:
        return jsonify({"sites": []})
    try:
        conn = _get_db_connection()
        cur = conn.cursor()
        sql = """
            SELECT hcsu.SITE_USE_ID, hcsu.LOCATION,
                   hl.ADDRESS1, hl.CITY, hl.STATE, hl.POSTAL_CODE,
                   hca.CUST_ACCOUNT_ID, hp.PARTY_NAME
              FROM HZ_CUST_ACCT_SITES_ALL hcas
              JOIN HZ_CUST_SITE_USES_ALL hcsu ON hcsu.CUST_ACCT_SITE_ID = hcas.CUST_ACCT_SITE_ID
              JOIN HZ_PARTY_SITES hps ON hps.PARTY_SITE_ID = hcas.PARTY_SITE_ID
              JOIN HZ_LOCATIONS hl ON hl.LOCATION_ID = hps.LOCATION_ID
              JOIN HZ_CUST_ACCOUNTS hca ON hca.CUST_ACCOUNT_ID = hcas.CUST_ACCOUNT_ID
              JOIN HZ_PARTIES hp ON hp.PARTY_ID = hca.PARTY_ID
             WHERE hcas.CUST_ACCOUNT_ID = :cust_account_id
               AND hcas.ORG_ID = :org_id
               AND hcsu.SITE_USE_CODE = 'SHIP_TO'
               AND hcsu.STATUS = 'A'
             ORDER BY hcsu.SITE_USE_ID
        """
        cur.execute(sql, cust_account_id=int(cust_account_id), org_id=int(ORG_ID))
        rows = cur.fetchall()
        cur.close()
        conn.close()
        sites = [{
            "ship_to_org_id": int(r[0]),
            "location": r[1] or "",
            "address": f"{r[2] or ''}, {r[3] or ''}, {r[4] or ''} {r[5] or ''}".strip(", "),
            "sold_to_org_id": int(r[6]),
            "customer_name": r[7] or ""
        } for r in rows]
        return jsonify({"sites": sites})
    except Exception as e:
        return jsonify({"sites": [], "error": str(e)})


@app.route("/api/lookup-prices", methods=["GET"])
def api_lookup_prices():
    """Get price lists for an inventory item."""
    item_id = request.args.get("item_id", "").strip()
    offset = int(request.args.get("offset", 0))
    if not item_id:
        return jsonify({"prices": [], "has_more": False})
    try:
        conn = _get_db_connection()
        cur = conn.cursor()
        sql = """
            SELECT * FROM (
                SELECT a.*, ROWNUM rnum FROM (
                    SELECT DISTINCT qlh.LIST_HEADER_ID AS price_list_id,
                                    qlt.NAME AS price_list_name,
                                    qll.OPERAND AS unit_list_price,
                                    qlh.START_DATE_ACTIVE
                      FROM QP_LIST_HEADERS_B qlh, QP_LIST_HEADERS_TL qlt, QP_LIST_LINES qll, QP_PRICING_ATTRIBUTES qpa
                     WHERE qlt.LIST_HEADER_ID = qlh.LIST_HEADER_ID AND qlt.LANGUAGE = USERENV('LANG')
                       AND qll.LIST_HEADER_ID = qlh.LIST_HEADER_ID
                       AND qpa.LIST_LINE_ID = qll.LIST_LINE_ID AND qpa.PRODUCT_ATTRIBUTE = 'PRICING_ATTRIBUTE1'
                       AND qpa.PRODUCT_ATTR_VALUE = :item_id_str
                       AND qlh.ACTIVE_FLAG = 'Y' AND qlh.LIST_TYPE_CODE = 'PRL' AND qlh.CURRENCY_CODE = 'USD'
                     ORDER BY qlh.START_DATE_ACTIVE DESC NULLS LAST
                ) a WHERE ROWNUM <= :upper_bound
            ) WHERE rnum > :offset
        """
        cur.execute(sql, item_id_str=str(item_id), upper_bound=offset + MAX_CHOICES, offset=offset)
        rows = cur.fetchall()
        cur.close()
        conn.close()
        prices = [{"price_list_id": int(r[0]), "price_list_name": r[1] or "", "unit_list_price": float(r[2])} for r in rows]
        return jsonify({"prices": prices, "has_more": len(prices) == MAX_CHOICES})
    except Exception as e:
        return jsonify({"prices": [], "has_more": False, "error": str(e)})


@app.route("/api/lookup-payment-terms", methods=["GET"])
def api_lookup_payment_terms():
    """Get available payment terms."""
    offset = int(request.args.get("offset", 0))
    try:
        conn = _get_db_connection()
        cur = conn.cursor()
        sql = """
            SELECT * FROM (
                SELECT a.*, ROWNUM rnum FROM (
                    SELECT DISTINCT TERM_ID, NAME
                    FROM RA_TERMS_VL
                    WHERE (START_DATE_ACTIVE IS NULL OR START_DATE_ACTIVE <= SYSDATE)
                      AND (END_DATE_ACTIVE IS NULL OR END_DATE_ACTIVE >= SYSDATE)
                    ORDER BY TERM_ID
                ) a WHERE ROWNUM <= :upper_bound
            ) WHERE rnum > :offset
        """
        cur.execute(sql, upper_bound=offset + MAX_CHOICES, offset=offset)
        rows = cur.fetchall()
        cur.close()
        conn.close()
        terms = [{"term_id": int(r[0]), "term_name": r[1] or ""} for r in rows]
        return jsonify({"terms": terms, "has_more": len(terms) == MAX_CHOICES})
    except Exception as e:
        return jsonify({"terms": [], "has_more": False, "error": str(e)})


@app.route("/api/lookup-salesreps", methods=["GET"])
def api_lookup_salesreps():
    """Get active sales representatives."""
    offset = int(request.args.get("offset", 0))
    try:
        conn = _get_db_connection()
        cur = conn.cursor()
        sql = """
            SELECT * FROM (
                SELECT a.*, ROWNUM rnum FROM (
                    SELECT DISTINCT jrs.SALESREP_ID, jrs.NAME AS salesrep_name
                    FROM JTF_RS_SALESREPS jrs
                     WHERE jrs.ORG_ID = :org_id AND jrs.STATUS = 'A' AND jrs.END_DATE_ACTIVE IS NULL
                     ORDER BY jrs.SALESREP_ID
                ) a WHERE ROWNUM <= :upper_bound
            ) WHERE rnum > :offset
        """
        cur.execute(sql, org_id=int(ORG_ID), upper_bound=offset + MAX_CHOICES, offset=offset)
        rows = cur.fetchall()
        cur.close()
        conn.close()
        reps = [{"salesrep_id": int(r[0]), "salesrep_name": r[1] or ""} for r in rows]
        return jsonify({"reps": reps, "has_more": len(reps) == MAX_CHOICES})
    except Exception as e:
        return jsonify({"reps": [], "has_more": False, "error": str(e)})


@app.route("/api/create-order", methods=["POST"])
def api_create_order():
    """Accept order form data and call PROCESS_ORDER REST service."""
    body = request.get_json(force=True)
    try:
        raw = create_order(body)
        parsed = parse_create_order_response(raw)
        _last_operation["type"] = "create"
        _last_operation["data"] = parsed
        _last_operation["raw"] = raw
        _last_operation["email_html"] = _generate_create_email_html(parsed)
        trigger_auto_email(f"New Order #{parsed.get('order_number', '')}")
        summary = generate_ai_summary("create", parsed)
        _last_operation["ai_summary"] = summary
        return jsonify({"reply": parsed, "type": "create_result", "ai_summary": summary})
    except requests.exceptions.HTTPError as e:
        return jsonify({"reply": f"REST API error: {e}", "type": "error"})
    except requests.exceptions.ConnectionError:
        return jsonify({
            "reply": "Could not connect to the Oracle EBS server. Please check network/VPN.",
            "type": "error",
        })
    except Exception as e:
        return jsonify({"reply": f"Something went wrong: {e}", "type": "error"})


@app.route("/api/orders-advanced", methods=["GET"])
def api_orders_advanced():
    """Get orders filtered by date range, year, or status."""
    range_type = request.args.get("range_type", "").strip() or None
    year = request.args.get("year", "").strip() or None
    status = request.args.get("status", "").strip() or None
    if year:
        year = int(year)
    try:
        data = get_orders_advanced(range_type=range_type, year=year, status=status)
        _last_operation["type"] = "orders_advanced"
        _last_operation["data"] = data
        _last_operation["raw"] = data
        _last_operation["email_html"] = _generate_orders_advanced_email_html(data, range_type, year, status)
        summary = generate_ai_summary("orders_advanced", data)
        _last_operation["ai_summary"] = summary
        return jsonify({"reply": data, "type": "orders_advanced", "ai_summary": summary})
    except Exception as e:
        return jsonify({"reply": f"Error querying orders: {e}", "type": "error"})


@app.route("/api/add-line", methods=["POST"])
def api_add_line():
    """Add a line item to an existing order."""
    body = request.get_json(force=True)
    header_id = body.get("header_id")
    if not header_id:
        return jsonify({"reply": "header_id is required.", "type": "error"})
    try:
        payload = add_line_payload(
            header_id=header_id,
            item_id=body.get("inventory_item_id"),
            ordered_item=body.get("ordered_item", ""),
            qty=body.get("ordered_quantity", 1),
            price=body.get("unit_list_price", 0),
            selling_price=body.get("unit_selling_price", 0),
            payment_term_id=body.get("payment_term_id", 4),
            price_list_id=body.get("price_list_id", 1000),
        )
        session = _build_session()
        resp = session.post(PROCESS_ORDER_URL, headers=HEADERS, data=json.dumps(payload), timeout=180)
        if resp.status_code in (200, 201, 204):
            try:
                raw = resp.json()
            except Exception:
                raw = {"raw": resp.text}
        else:
            resp.raise_for_status()
        parsed = parse_create_order_response(raw)
        summary = generate_ai_summary("add_line", parsed)
        return jsonify({"reply": parsed, "type": "add_line_result", "ai_summary": summary})
    except requests.exceptions.HTTPError as e:
        return jsonify({"reply": f"REST API error: {e}", "type": "error"})
    except requests.exceptions.ConnectionError:
        return jsonify({"reply": "Could not connect to the Oracle EBS server.", "type": "error"})
    except Exception as e:
        return jsonify({"reply": f"Something went wrong: {e}", "type": "error"})


@app.route("/api/send-email", methods=["POST"])
def api_send_email():
    """Send the last operation result as an email."""
    body = request.get_json(force=True)
    to_email = (body.get("to_email") or "").strip()
    subject = (body.get("subject") or "Oracle ERP Order Details Update").strip()
    body_text = (body.get("body") or "Please find the requested order details structured below.").strip()

    if not to_email:
        return jsonify({"reply": "Email address is required.", "type": "error"})

    html_payload = _last_operation.get("email_html", "")
    if not html_payload:
        return jsonify({"reply": "No recent operation data to send. Please look up an order first.", "type": "error"})

    success, msg = send_erp_email(to_email, subject, body_text, html_payload)
    if success:
        return jsonify({"reply": f"Email sent successfully to {to_email}!", "type": "email_success"})
    else:
        return jsonify({"reply": f"Error sending email: {msg}", "type": "error"})


@app.route("/api/download-excel")
def api_download_excel():
    """Generate and return an Excel file of the last operation data."""
    if not OPENPYXL_AVAILABLE:
        return jsonify({"error": "openpyxl not installed. Run: pip install openpyxl"}), 500

    if not _last_operation.get("type") or not _last_operation.get("data"):
        return jsonify({"error": "No operation data available. Please look up an order first."}), 400

    excel_bytes = _generate_excel()
    if not excel_bytes:
        return jsonify({"error": "Excel generation failed."}), 500

    fname = _get_excel_filename()
    return Response(
        excel_bytes,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={fname}"},
    )


@app.route("/api/email-preview")
def api_email_preview():
    """Return the current email HTML preview."""
    return jsonify({
        "html": _last_operation.get("email_html", ""),
        "has_data": _last_operation["type"] is not None,
        "operation_type": _last_operation["type"],
    })


@app.route("/api/auto-email", methods=["POST"])
def api_auto_email():
    """Toggle auto-email after every operation."""
    data = request.get_json(force=True)
    _auto_email["enabled"] = bool(data.get("enabled", False))
    _auto_email["to"] = (data.get("to") or "").strip()
    return jsonify({"ok": True, **_auto_email})


@app.route("/api/auto-email", methods=["GET"])
def api_auto_email_status():
    """Get auto-email config."""
    return jsonify(_auto_email)


@app.route("/api/scheduler/start", methods=["POST"])
def api_scheduler_start():
    """Start the recurring email scheduler."""
    data = request.get_json(force=True)
    to_email = (data.get("to_email") or "").strip()
    customer_name = (data.get("customer_name") or "").strip()
    interval_str = (data.get("interval") or "60").strip()

    if not to_email:
        return jsonify({"error": "Recipient email is required."}), 400

    # Parse interval
    mins = 60
    interval_lower = interval_str.lower()
    if "daily" in interval_lower or "day" in interval_lower:
        mins = 1440
    elif "hour" in interval_lower:
        nums = re.findall(r"\d+", interval_str)
        mins = int(nums[0]) * 60 if nums else 60
    elif "min" in interval_lower:
        nums = re.findall(r"\d+", interval_str)
        mins = int(nums[0]) if nums else 60
    else:
        nums = re.findall(r"\d+", interval_str)
        if nums:
            mins = int(nums[0])

    _start_scheduler(mins, customer_name, to_email)
    return jsonify({
        "ok": True,
        "message": f"Scheduler started: every {mins} min(s) → {to_email}",
        **_scheduler,
    })


@app.route("/api/scheduler/stop", methods=["POST"])
def api_scheduler_stop():
    """Stop the recurring email scheduler."""
    _stop_scheduler()
    return jsonify({"ok": True, "message": "Scheduler stopped.", **_scheduler})


@app.route("/api/scheduler/status")
def api_scheduler_status():
    """Get the current scheduler state."""
    return jsonify(_scheduler)


# ── Auto-Loop Endpoints ──────────────────────────────────────
@app.route("/api/auto-loop/start", methods=["POST"])
def api_auto_loop_start():
    """Start auto email loop: fetch all customers, send email for each, repeat."""
    body = request.get_json(force=True)
    to_email = (body.get("to_email") or "").strip()
    if not to_email:
        return jsonify({"error": "Recipient email is required."}), 400

    _start_auto_loop(to_email)
    return jsonify({"ok": True, "message": f"Auto-loop started -> {to_email}", **_auto_loop})


@app.route("/api/auto-loop/stop", methods=["POST"])
def api_auto_loop_stop():
    """Stop the auto email loop."""
    _stop_auto_loop()
    return jsonify({"ok": True, "message": "Auto-loop stopped.", **_auto_loop})


@app.route("/api/auto-loop/status")
def api_auto_loop_status():
    """Get auto-loop status."""
    return jsonify(_auto_loop)


@app.route("/api/voice", methods=["POST"])
def api_voice():
    """Server-side voice recognition using Vosk (fallback when browser Speech API unavailable)."""
    if not VOSK_AVAILABLE:
        return jsonify({"text": "", "error": "Voice packages not installed. Use browser microphone instead."})
    try:
        text = recognize_speech(VOICE_DURATION)
        if text:
            return jsonify({"text": text})
        return jsonify({"text": "", "error": "No speech detected. Try again."})
    except Exception as e:
        return jsonify({"text": "", "error": f"Voice error: {e}"})


@app.route("/api/speak", methods=["POST"])
def api_speak():
    """Server-side TTS using pyttsx3 (optional)."""
    if not PYTTSX3_AVAILABLE:
        return jsonify({"ok": False, "error": "pyttsx3 not installed."})
    data = request.get_json(force=True)
    text = (data.get("text") or "").strip()
    if not text:
        return jsonify({"ok": False, "error": "No text provided."})
    t = threading.Thread(target=speak, args=(text,), daemon=True)
    t.start()
    return jsonify({"ok": True})


if __name__ == "__main__":
    print("Sales Order Chatbot running at http://localhost:5000")
    app.run(host="0.0.0.0", port=5000, debug=True, use_reloader=False)
